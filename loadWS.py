#!/usr/bin/env python

import MySQLdb as mdb
from openpyxl import load_workbook
import sys

import time
from datetime import datetime
import uuid

import pprint

pp = pprint.PrettyPrinter(indent = 1)

def loadWStoEMED(type, wsdata, dbh):

	# Get a curstor into the database
	cur = dbh.cursor(mdb.cursors.DictCursor)
	
	wb = load_workbook(filename = wsdata[type]['fname'])
	ws = wb[wsdata[type]['input_sheetname']]
	
	skip_header = False
	if type == 'App' and ws["A1"].value == "AppName":
		skip_header = True
	if type == 'Trap' and ws["A1"].value == "EventName":
		skip_header = True
	
	if type == 'App':
		sql_preamble = 'INSERT INTO `eventsApp` ( AppName, EventName, AlertMessage, \
			ThresholdValue, ThresholdUnit, EventSeverity, EventID, EventMessage, \
			AlertName, ThresholdName, EventGUID, AppGUID, AlertGUID, ThresholdGUID \
			) VALUES ('
	
	if type == 'Trap':
		sql_preamble = 'INSERT INTO `eventsTrap` ( 	`EventName`, `EventSeverity`, \
			`EventMessage`,	`EventGUID`, `PowerpackGUID` \
			) VALUES ('
	
	sql_epilogue = ' );'
		
	rowcount = 1
	for row in ws.rows:
		sql_values = ''
		#print "Row %d" % (rowcount)
		for cell in row:
			if skip_header == True and rowcount != 1:
				#print sql_values
				#sql_values = sql_values + ' \'' + cell.value + '\','
				sql_values = '%s \'%s\',' % (sql_values, cell.value)
				#print(cell.value)
		rowcount = rowcount + 1
		sql_values = sql_values[:-1]
		if skip_header == False or rowcount >= 3:
			sql_statement = sql_preamble + sql_values + sql_epilogue
			#print sql_statement	
			try:
				cur.execute(sql_statement)
			except e:
				print "Error %d: %s" % (e.args[0], e.args[1])
				dbh.rollback
	
	dbh.commit()
	cur.close()
	